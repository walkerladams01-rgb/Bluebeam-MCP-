"""Bluebeam markup parsing (against tests/bb_synth.py ground truth) and the read-only tools."""
import asyncio
import csv
import json
import time

import pymupdf
import pytest
from mcp.server.fastmcp import FastMCP

from bluebeam import markups_read as mr
from bluebeam import tools_read as tr
from bluebeam.core import DocumentError
from tests import bb_synth as S

LISTED = {"area", "length", "count", "callout", "cloud", "text", "rectangle"}     # kinds parse_markups lists


# =============================================================================
# fixtures and helpers
# =============================================================================

@pytest.fixture(scope="module", params=[True, False], ids=["indirect-measure", "direct-measure"])
def synth(request, tmp_path_factory):
    path = str(tmp_path_factory.mktemp("mr") / "synth.pdf")
    return path, S.build_takeoff_pdf(path, indirect_measure=request.param)


@pytest.fixture(scope="module")
def synth1(tmp_path_factory):
    """One synth PDF (indirect measure) for tests that do not need both variants."""
    path = str(tmp_path_factory.mktemp("mr1") / "synth1.pdf")
    return path, S.build_takeoff_pdf(path)


def _by_id(mks):
    return {m.id: m for m in mks}


def _mcp():
    mcp = FastMCP("t")
    tr.register_read_tools(mcp)
    return mcp


def _call(name, **kw):
    out = asyncio.run(_mcp().call_tool(name, kw))
    content = out[0] if isinstance(out, tuple) else out
    assert len(content) == 1 and content[0].type == "text"
    return json.loads(content[0].text)


def _pdf(tmp_path, name="t.pdf", pages=1, width=612, height=792):
    doc = pymupdf.open()
    for _ in range(pages):
        doc.new_page(width=width, height=height)
    return doc, str(tmp_path / name)


def _save(doc, path):
    doc.save(path)
    doc.close()
    return path


# =============================================================================
# the PDF object reader and small helpers
# =============================================================================

def test_reader_strings_names_numbers_refs():
    d = mr.parse_pdf_object(
        r"<</A(a\(b\)c\101 \\ \n)/B<FEFF00480069>/C[1 0 R .5 -.25 7 true false null /N#20x]"
        r"/D<</E 3 0 R/F(nested (parens) ok)>> %comment\n/G 9>>")
    assert d["A"] == "a(b)cA \\ \n"
    assert d["B"] == "Hi"
    assert d["C"] == [1, 0.5, -0.25, 7, True, False, None, "N x"]
    assert isinstance(d["C"][0], mr.Ref) and int(d["C"][0]) == 1
    assert isinstance(d["D"]["E"], mr.Ref) and d["D"]["F"] == "nested (parens) ok"
    assert "G" not in d                                                # swallowed by the comment (no newline)
    assert mr.parse_pdf_object("<</G 9>>") == {"G": 9}


def test_reader_text_encodings_and_junk():
    assert mr.decode_pdf_text(b"\xfe\xff\x00\xb0\x00A") == "\u00b0A"
    assert mr.decode_pdf_text(b"\xff\xfe\xb0\x00A\x00") == "\u00b0A"
    assert mr.decode_pdf_text(b"caf\xc3\xa9") == "caf\u00e9"                # UTF-8
    assert mr.decode_pdf_text(b"\xb0") == "\u00b0"                          # PDFDocEncoding fallback
    assert mr.parse_pdf_object("<</U<B0>>>") == {"U": "\u00b0"}             # 1-byte hex string (angle unit)
    for junk in ("", "<<", "<</A", "[1 2", "(unterminated", ">> ]] ))", "<</A /B /C>>", "<</A <<>> >> >>"):
        mr.parse_pdf_object(junk)                                          # must terminate, never raise


def test_iso_date_color_and_rich_text():
    assert mr.iso_date("D:20260928120000-07'00'") == "2026-09-28T12:00:00-07:00"
    assert mr.iso_date("D:20260928120000Z") == "2026-09-28T12:00:00Z"
    assert mr.iso_date("D:20260928") == "2026-09-28T00:00:00"
    assert mr.iso_date("D:20261399") == "D:20261399"                       # impossible date kept as written
    assert mr.iso_date("garbage") == "garbage" and mr.iso_date("") is None
    assert mr._hex_color([1, 0, 0]) == "#FF0000"
    assert mr._hex_color([0.5]) == "#808080"
    assert mr._hex_color([0, 0, 0, 1]) == "#000000" and mr._hex_color([0, 0, 0, 0]) == "#FFFFFF"
    assert mr._hex_color([]) is None and mr._hex_color(["x"]) is None
    rc = ('<?xml version="1.0"?><body xmlns="http://www.w3.org/1999/xhtml"><style>p{x:y}</style>'
          "<p>Line <b>one</b> &amp; two</p><p>Line 3</p></body>")
    assert mr.plain_rich_text(rc) == "Line one & two\nLine 3"
    assert mr.plain_rich_text("") == ""


# =============================================================================
# parse_markups vs ground truth
# =============================================================================

def test_parse_matches_ground_truth(synth):
    path, exp = synth
    with pymupdf.open(path) as doc:
        mks = _by_id(mr.parse_markups(doc))
    expected = {k: v for k, v in exp["markups"].items() if v["kind"] in LISTED}
    assert set(mks) == set(expected), "replies and statuses are not listed by default"
    for nm, e in expected.items():
        m = mks[nm]
        assert (m.xref, m.page, m.pdf_type, m.intent) == (e["xref"], e["page"], e["pdf_type"], e["intent"])
        assert m.author == e["author"] and m.layer == e["layer"]
        assert m.page_label == exp["page_labels"][e["page"]]
        assert m.group_parent_id == e["group_parent"]
        assert m.custom_columns == e["custom_columns"]
        assert m.created == m.modified == "2026-09-28T12:00:00-07:00"
        if e["subject"] is not None:
            assert m.subject == e["subject"]
        if e["vertex_count"]:
            assert len(m.vertices) == e["vertex_count"]
        me = m.measurement
        if e["kind"] in ("area", "length", "count"):
            assert me["kind"] == e["kind"] and me["unit"] == e["unit"] and me["formatted"] == e["formatted"]
            assert me["value"] == pytest.approx(e["value"], abs=0.01)
            assert me["count"] == e["count"]
            assert me["depth"] == e["depth"] and me["depth_unit"] == e["depth_unit"]
            assert me["volume_unit"] == e["volume_unit"]
            if e["volume"] is None:
                assert me["volume"] is None
            else:
                assert me["volume"] == pytest.approx(e["volume"], abs=0.01)
        else:
            assert me is None
            if e["formatted"]:
                assert m.contents == e["formatted"]


def test_measurement_details_scale_and_check_value(synth):
    path, exp = synth
    with pymupdf.open(path) as doc:
        mks = _by_id(mr.parse_markups(doc))
    slab = next(m for m in mks.values() if m.subject == "Concrete Slab")
    me = slab.measurement
    assert me["scale_text"] == "1 in = 30 ft' in\"" and me["units_per_point"] == 0.4166667
    assert me["x_unit"] == "ft" and me["unit_factor"] == 1.0
    assert me["computed"] == pytest.approx(3333.33, abs=0.01) and me["contents_value"] == 3333.33
    line = next(m for m in mks.values() if m.pdf_type == "Line")
    assert line.measurement["unit"] == "ft" and line.measurement["formatted"] == "61'-11 1/4\""
    assert line.measurement["contents_value"] == pytest.approx(61.9375)
    assert line.measurement["kind"] == "length"
    peri = next(m for m in mks.values() if m.pdf_type == "PolyLine")
    assert peri.measurement["kind"] == "length"                             # an open polyline: no closing segment
    assert peri.measurement["computed"] == pytest.approx(exp["markups"][peri.id]["value"], abs=1e-3)


def test_groups_counts_and_cloud(synth):
    path, exp = synth
    with pymupdf.open(path) as doc:
        mks = _by_id(mr.parse_markups(doc))
    for parent, kids in exp["groups"].items():
        for kid in kids:
            assert mks[kid].group_parent_id == parent
        assert mks[parent].group_parent_id is None
    count_parent = next(p for p, kids in exp["groups"].items() if exp["markups"][p]["kind"] == "count")
    members = exp["groups"][count_parent]
    assert len(members) == 2
    assert all(mks[k].measurement["member_of"] == count_parent for k in members)
    assert "member_of" not in mks[count_parent].measurement
    assert {mks[k].measurement["count"] for k in [count_parent, *members]} == {3}       # N on every member
    assert mks[count_parent].measurement["count_style"] == "Checkmark"
    cloud = next(m for m in mks.values() if m.intent == "PolygonCloud")
    assert cloud.cloudy and cloud.measurement is None
    callout = mks[cloud.group_parent_id]
    assert callout.intent == "FreeTextCallout" and not callout.cloudy


def test_replies_and_status(synth):
    path, exp = synth
    with pymupdf.open(path) as doc:
        plain = _by_id(mr.parse_markups(doc, pages=[2]))
        with_replies = _by_id(mr.parse_markups(doc, pages=[2], include_replies=True))
    (parent, reply_ids), = exp["replies"].items()
    assert reply_ids[0] not in plain
    reply = with_replies[reply_ids[0]]
    assert reply.is_reply and reply.reply_to_id == parent and reply.contents == "Checked against the detail"
    assert reply.author == S.REVIEWER and reply.group_parent_id is None
    for nm, st in exp["status"].items():
        for coll in (plain, with_replies):
            assert coll[nm].status["state"] == st["state"]
            assert coll[nm].status["model"] == st["model"] and coll[nm].status["by"] == st["by"]
            assert coll[nm].status["date"] == "2026-09-28T12:00:00-07:00"
        assert st["id"] not in with_replies, "status entries are never listed"


def test_latest_status_wins(tmp_path):
    path = str(tmp_path / "st.pdf")
    exp = S.build_takeoff_pdf(path)
    (parent, st), = [(k, v) for k, v in exp["status"].items() if v["state"] == "Rejected"]
    with pymupdf.open(path) as doc:
        page = doc[2]
        later = S.add_status(doc, page, exp["markups"][parent]["xref"], "Completed", model="Review")
        doc.xref_set_key(later, "CreationDate", "(D:20261001120000Z)")
        doc.saveIncr()
    with pymupdf.open(path) as doc:
        m = _by_id(mr.parse_markups(doc, pages=[2]))[parent]
    assert m.status["state"] == "Completed"


def test_pages_argument_and_page_labels(synth1):
    path, exp = synth1
    with pymupdf.open(path) as doc:
        assert mr.parse_markups(doc, pages=[0]) == []
        p3 = mr.parse_markups(doc, pages=[2])
        assert {m.page for m in p3} == {3}
        assert {m.page_label for m in p3} == {exp["page_labels"][3]}
        assert len(mr.parse_markups(doc)) == 12
        assert [m.page for m in mr.parse_markups(doc, pages=[2, 1, 1])] == [2] * 8 + [3] * 4      # sorted, unique


def test_to_dict_summary_and_full(synth1):
    path, exp = synth1
    with pymupdf.open(path) as doc:
        mks = _by_id(mr.parse_markups(doc))
    riprap = next(m for m in mks.values() if m.subject == "Riprap Area" and m.custom_columns)
    s = riprap.to_dict()
    assert s["type"] == "Polygon" and s["page"] == 2 and s["layer"] == "Takeoff"
    assert s["measurement"] == {"kind": "area", "value": 5208.33, "unit": "sf", "formatted": "5,208.33 sf",
                                "volume": 96.45, "volume_unit": "cu yd"}
    assert s["custom_columns"]["Bid Item"] == "RIPRAP-12"
    assert "rect" not in s and "vertices" not in s and "locked" not in s
    json.dumps(s)
    f = riprap.to_dict("full")
    assert f["rect"] == [99.0, 99.0, 301.0, 251.0] and f["vertex_count"] == 4 and len(f["vertices"]) == 4
    assert f["color"] == "#FF0000" and f["fill_color"] == "#FFCCCC" and f["opacity"] == 0.5
    assert f["line_width"] == 1.0 and f["intent"] == "PolygonDimension"
    assert f["measurement"]["units_per_point"] == 0.4166667 and f["measurement"]["depth"] == 6.0
    json.dumps(f)
    text_box = next(m for m in mks.values() if m.subject == "Text Box")
    assert text_box.rich_text_plain == "General note"
    assert "measurement" not in text_box.to_dict()
    assert "rich_text_plain" not in text_box.to_dict("full")                # identical to contents: not repeated


def test_summary_contents_truncated_and_status(synth1):
    path, _ = synth1
    with pymupdf.open(path) as doc:
        mks = mr.parse_markups(doc, pages=[2])
    callout = next(m for m in mks if m.intent == "FreeTextCallout")
    assert callout.to_dict()["status"] == "Accepted"
    callout.contents = "x" * 500
    assert len(callout.to_dict()["contents"]) == 300 and callout.to_dict()["contents"].endswith("\u2026")


# =============================================================================
# measure, scales, layers, custom columns
# =============================================================================

def test_parse_measure_indirect_and_direct(synth):
    path, exp = synth
    with pymupdf.open(path) as doc:
        mks = mr.parse_markups(doc)
        for m in mks:
            pm = mr.parse_measure(doc, m.xref)
            if m.intent in ("PolygonDimension", "PolyLineDimension", "LineDimension"):
                assert pm["scale_text"] == "1 in = 30 ft' in\"" and pm["units_per_point"] == 0.4166667
                assert pm["x_unit"] == "ft" and pm["distance"] == {"unit": "ft", "factor": 1.0}
                assert pm["area"] == {"unit": "sf", "factor": 1.0}
                assert pm["volume"] == {"unit": "cu yd", "factor": pytest.approx(0.03703704)}
                assert pm["angle"] == {"unit": "\u00b0", "factor": 1.0}             # <B0> hex unit
                assert exp["markups"][m.id]["measure_indirect"] in (True, False)
            elif m.intent is None:
                assert pm is None


def test_measure_indirect_flag_really_differs(synth):
    path, exp = synth
    with pymupdf.open(path) as doc:
        for nm, e in exp["markups"].items():
            if e["kind"] in ("area", "length"):
                kind = doc.xref_get_key(e["xref"], "Measure")[0]
                assert kind == ("xref" if e["measure_indirect"] else "dict")


def test_page_scales(synth1):
    path, _ = synth1
    with pymupdf.open(path) as doc:
        assert mr.page_scales(doc, 0) == []
        got = mr.page_scales(doc, 1)
        assert [s["source"] for s in got] == ["viewport", "markup"]
        assert got[0]["scale"] == "1 in = 30 ft' in\"" and got[0]["units_per_point"] == 0.4166667
        assert got[0]["bbox"] == [0.0, 0.0, 612.0, 792.0] and got[1]["markups"] == 5
        assert mr.page_scales(doc, 1, mr.parse_markups(doc)) == got           # passing parsed markups is equivalent
        assert mr.page_scales(doc, 2) == []


def test_list_layers(synth1):
    path, _ = synth1
    with pymupdf.open(path) as doc:
        layers = mr.list_layers(doc)
        assert [(l["name"], l["visible"], l["locked"], l["markup_count"]) for l in layers] == \
            [("Takeoff", True, False, 8), ("Review", True, False, 4)]
        review = layers[1]["xref"]
        doc.set_layer(-1, on=[layers[0]["xref"]], off=[review])
        assert [l["visible"] for l in mr.list_layers(doc)] == [True, False]
        doc.xref_set_key(doc.pdf_catalog(), "OCProperties/D/Locked", f"[{review} 0 R]")
        assert [l["locked"] for l in mr.list_layers(doc)] == [False, True]


def test_list_layers_none(sample_pdf):
    with pymupdf.open(sample_pdf) as doc:
        assert mr.list_layers(doc) == []


def test_custom_column_defs_and_absence(synth1, tmp_path):
    path, exp = synth1
    with pymupdf.open(path) as doc:
        defs = mr.custom_column_defs(doc)
    assert [d["name"] for d in defs] == exp["custom_column_defs"]
    assert [d["subtype"] for d in defs] == ["Text", "Number", "Choice"]
    plain = str(tmp_path / "nocols.pdf")
    S.build_takeoff_pdf(plain, custom_columns=False)
    with pymupdf.open(plain) as doc:
        assert mr.custom_column_defs(doc) == []
        assert all(m.custom_columns == {} for m in mr.parse_markups(doc))


def test_custom_columns_survive_bad_structures(tmp_path):
    path = str(tmp_path / "badcols.pdf")
    S.build_takeoff_pdf(path)
    with pymupdf.open(path) as doc:
        doc.xref_set_key(doc.pdf_catalog(), "BSIColumnData", "42")            # not an array
        assert mr.custom_column_defs(doc) == []
        cols = [m.custom_columns for m in mr.parse_markups(doc) if m.custom_columns]
        assert cols and set(cols[0]) == {"Column 1", "Column 2", "Column 3"}    # positional names, no crash
        doc.xref_set_key(doc.pdf_catalog(), "BSIColumnData", "[1 2 3]")         # array of non-dicts
        assert mr.custom_column_defs(doc) == []
        doc.xref_set_key(doc.pdf_catalog(), "BSIColumnData", "null")
        mr.parse_markups(doc)


def test_raw_keys(synth):
    path, exp = synth
    e = next(v for v in exp["markups"].values() if v["subject"] == "Concrete Slab")
    with pymupdf.open(path) as doc:
        keys = mr.raw_keys(doc, e["xref"])
        short = mr.raw_keys(doc, e["xref"], max_len=10)
    assert keys["IT"] == "PolygonDimension" and keys["MeasurementTypes"] == "129" and keys["Subj"] == "Concrete Slab"
    assert keys["Measure"].startswith("<<") != e["measure_indirect"]
    assert "Vertices" in keys and len(short["DS"]) <= 10 and short["DS"].endswith("…")


# =============================================================================
# hand-made Bluebeam quirks
# =============================================================================

def _poly(page, pts, closed=True):
    return (page.add_polygon_annot if closed else page.add_polyline_annot)([pymupdf.Point(*p) for p in pts])


def test_angle_hex_unit_and_utf16_contents(tmp_path):
    doc, path = _pdf(tmp_path)
    page = doc[0]
    a = _poly(page, [(100, 100), (200, 100), (200, 200)], closed=False)
    x = a.xref
    doc.xref_set_key(x, "NM", S.pdf_str("angle-1"))
    doc.xref_set_key(x, "IT", "/PolyLineAngle")
    doc.xref_set_key(x, "MeasurementTypes", "1152")
    doc.xref_set_key(x, "Contents", S.utf16_hex("90.00\u00b0"))                # UTF-16BE with BOM, like CAD/Revu
    S.attach_measure(doc, x, S.measure_dict(30), indirect=False)
    _save(doc, path)
    with pymupdf.open(path) as doc:
        (m,) = mr.parse_markups(doc)
    assert m.contents == "90.00\u00b0"
    assert m.measurement["kind"] == "angle" and m.measurement["unit"] == "\u00b0"
    assert m.measurement["value"] == pytest.approx(90) and m.measurement["contents_value"] == 90.0
    assert m.measurement["formatted"] == "90.00\u00b0"


def test_unscaled_measurement_uses_bluebeams_text(tmp_path):
    doc, path = _pdf(tmp_path)
    a = _poly(doc[0], [(0, 0), (100, 0), (100, 100), (0, 100)])
    doc.xref_set_key(a.xref, "IT", "/PolygonDimension")
    doc.xref_set_key(a.xref, "MeasurementTypes", "129")
    doc.xref_set_key(a.xref, "Contents", "(100 sf)")
    _save(doc, path)
    with pymupdf.open(path) as doc:
        (m,) = mr.parse_markups(doc)
    me = m.measurement
    assert me["kind"] == "area" and me["units_per_point"] is None and me["computed"] is None
    assert me["value"] == 100.0 and me["unit"] == "sf"


def test_garbage_measure_and_contents_do_not_crash(tmp_path):
    doc, path = _pdf(tmp_path)
    a = _poly(doc[0], [(0, 0), (100, 0), (100, 100)])
    doc.xref_set_key(a.xref, "IT", "/PolygonDimension")
    doc.xref_set_key(a.xref, "Measure", "5")                                   # not a dictionary
    doc.xref_set_key(a.xref, "Contents", "42")                                 # a number, not a string
    b = _poly(doc[0], [(10, 10), (50, 10), (50, 50)])
    doc.xref_set_key(b.xref, "IT", "/PolygonDimension")
    doc.xref_set_key(b.xref, "Measure", "<</Type/Measure/X[]>>")               # no scale entries
    doc.xref_set_key(b.xref, "Depth", "(deep)")
    _save(doc, path)
    with pymupdf.open(path) as doc:
        ms = mr.parse_markups(doc)
    assert len(ms) == 2 and all(m.measurement["computed"] is None for m in ms)
    assert all(m.contents == "" for m in ms)


def _raw_string_pdf(path, contents: bytes, subject: bytes) -> str:
    """A minimal hand-written PDF whose one annotation carries raw string bytes (as Revu writes them)."""
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Annots [4 0 R] >>",
            b"<< /Type /Annot /Subtype /Square /Rect [10 10 50 50] /NM (x1) /Contents (" + contents
            + b") /Subj (" + subject + b") /T <FEFF00E900410042> >>"]
    out, offsets = bytearray(b"%PDF-1.7\n"), []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    with open(path, "wb") as f:
        f.write(out)
    return str(path)


@pytest.mark.parametrize("contents,subject", [
    (b"caf\xe9 45\xb0", b"Riprap \xb0"),                                    # PDFDocEncoding / Latin-1 bytes
    ("caf\u00e9 45\u00b0".encode("utf-8"), "Riprap \u00b0".encode("utf-8")),   # UTF-8 bytes
    (b"\xfe\xff\x00c\x00a\x00f\x00\xe9\x00 \x004\x005\x00\xb0",
     b"\xfe\xff\x00R\x00i\x00p\x00r\x00a\x00p\x00 \x00\xb0"),               # UTF-16BE, raw bytes
    (rb"\376\377\000c\000a\000f\000\351\000 \0004\0005\000\260",
     rb"\376\377\000R\000i\000p\000r\000a\000p\000 \000\260"),              # UTF-16BE, octal escapes
])
def test_string_encodings_from_real_bytes(tmp_path, contents, subject):
    path = _raw_string_pdf(tmp_path / "raw.pdf", contents, subject)
    with pymupdf.open(path) as doc:
        (m,) = mr.parse_markups(doc)
    assert m.contents == "caf\u00e9 45\u00b0"
    assert m.subject == "Riprap \u00b0" and m.author == "\u00e9AB"


def test_flags_locked_hidden(tmp_path):
    doc, path = _pdf(tmp_path)
    page = doc[0]
    locked = page.add_rect_annot(pymupdf.Rect(10, 10, 50, 50))
    hidden = page.add_rect_annot(pymupdf.Rect(60, 10, 100, 50))
    normal = page.add_rect_annot(pymupdf.Rect(110, 10, 150, 50))
    doc.xref_set_key(locked.xref, "F", "132")                                  # Print + Locked
    doc.xref_set_key(hidden.xref, "F", "2")
    _save(doc, path)
    with pymupdf.open(path) as doc:
        ms = {m.xref: m for m in mr.parse_markups(doc)}
    assert ms[locked.xref].locked and not ms[locked.xref].hidden
    assert ms[hidden.xref].hidden and not ms[hidden.xref].locked
    assert not ms[normal.xref].locked and not ms[normal.xref].hidden
    assert ms[locked.xref].to_dict()["locked"] is True and "hidden" not in ms[locked.xref].to_dict()


def test_markup_without_nm_gets_xref_id_and_ocmd_layer(tmp_path):
    doc, path = _pdf(tmp_path)
    page = doc[0]
    a = page.add_rect_annot(pymupdf.Rect(10, 10, 50, 50))
    doc.xref_set_key(a.xref, "NM", "null")
    g1 = S.add_layer(doc, "Layer One")
    g2 = S.add_layer(doc, "Layer Two")
    ocmd = doc.get_new_xref()
    doc.update_object(ocmd, f"<</Type/OCMD/OCGs[{g1} 0 R {g2} 0 R]>>")
    doc.xref_set_key(a.xref, "OC", f"{ocmd} 0 R")
    _save(doc, path)
    with pymupdf.open(path) as doc:
        (m,) = mr.parse_markups(doc)
    assert m.id == f"xref:{a.xref}" and m.layer == "Layer One, Layer Two"


def test_rotated_page_and_ink_do_not_break(tmp_path):
    doc, path = _pdf(tmp_path)
    page = doc[0]
    page.add_ink_annot([[(10, 10), (20, 30), (40, 20)]])
    page.add_line_annot((10, 10), (50, 60))
    page.set_rotation(90)
    _save(doc, path)
    with pymupdf.open(path) as doc:
        ms = mr.parse_markups(doc)
    ink = next(m for m in ms if m.pdf_type == "Ink")
    assert ink.vertices is None and len(ms) == 2


# =============================================================================
# tools: list / get / takeoff
# =============================================================================

def test_list_markups_default_and_paging(synth1):
    path, _ = synth1
    r = tr.list_markups(path)
    assert r["file"] == "synth1.pdf" and r["total"] == r["returned"] == 12 and "truncated" not in r
    assert len(tr.list_markups(path, include_replies=True)["markups"]) == 13
    assert tr.list_markups(path, pages="2")["total"] == 8
    assert tr.list_markups(path, pages=[3])["total"] == 4
    assert tr.list_markups(path, pages="1")["total"] == 0
    with pytest.raises(DocumentError, match="out of range"):
        tr.list_markups(path, pages="9")
    small = tr.list_markups(path, limit=5)
    assert small["returned"] == 5 and small["total"] == 12 and small["truncated"] and "narrow" in small["note"]


@pytest.mark.parametrize("kwargs,expected", [
    ({"types": ["area"]}, 3), ({"types": "length,count"}, 5), ({"types": ["Polygon"]}, 7),
    ({"types": ["measurement"]}, 8), ({"types": ["PolygonCount"]}, 3), ({"types": ["freetext"]}, 2),
    ({"subject": "riprap"}, 2), ({"subject": "CLOUD"}, 2), ({"author": "estimator"}, 12),
    ({"layer": "review"}, 4), ({"layer": "take", "types": ["area"]}, 3), ({"subject": "nothing"}, 0),
])
def test_list_markups_filters(synth1, kwargs, expected):
    assert tr.list_markups(synth1[0], **kwargs)["total"] == expected


def test_list_markups_reply_filters_and_detail(synth1):
    path, _ = synth1
    r = tr.list_markups(path, author="reviewer", include_replies=True)
    assert r["total"] == 1 and r["markups"][0]["reply_to_id"]
    full = tr.list_markups(path, pages=[2], detail="full", types=["area"])["markups"]
    assert all("rect" in m and "vertices" in m for m in full)
    with pytest.raises(DocumentError, match="detail"):
        tr.list_markups(path, detail="everything")


def test_get_markup(synth1):
    path, exp = synth1
    callout = next(k for k, v in exp["markups"].items() if v["kind"] == "callout")
    cloud = exp["groups"][callout][0]
    reply = exp["replies"][callout][0]
    r = tr.get_markup(path, callout)
    assert r["markup"]["intent"] == "FreeTextCallout" and r["markup"]["status"]["state"] == "Accepted"
    assert r["group_children"] == [cloud]
    assert [x["id"] for x in r["replies"]] == [reply] and r["replies"][0]["contents"] == "Checked against the detail"
    assert r["raw_keys"]["IT"] == "FreeTextCallout" and "CL" in r["raw_keys"]
    c = tr.get_markup(path, cloud)
    assert c["markup"]["group_parent_id"] == callout and c["markup"]["cloudy"] is True
    assert tr.get_markup(path, f"xref:{exp['markups'][cloud]['xref']}")["markup"]["id"] == cloud
    assert tr.get_markup(path, str(exp["markups"][cloud]["xref"]))["markup"]["id"] == cloud
    assert tr.get_markup(path, reply)["markup"]["reply_to_id"] == callout
    with pytest.raises(DocumentError, match="Markup not found"):
        tr.get_markup(path, "no-such-id")
    status_id = exp["status"][callout]["id"]
    with pytest.raises(DocumentError, match="review-status"):
        tr.get_markup(path, status_id)


def test_takeoff_summary_tool(synth1):
    path, exp = synth1
    r = tr.takeoff_summary(path)
    assert r["group_by"] == ["subject"] and r["file"] == "synth1.pdf"
    areas = {(t["kind"], t["unit"]): t["total"] for t in r["totals"]}
    assert areas[("area", "sf")] == pytest.approx(exp["totals"]["area_sf"], abs=0.01)
    assert areas[("count", "count")] == 3
    assert all("markups" in t and "count" not in t for t in r["totals"])
    assert [t["volume_unverified"] for t in r["totals"] if t["kind"] == "volume"] == [True] and "notes" in r
    by_layer = tr.takeoff_summary(path, group_by=["layer", "kind"], pages="2")
    assert {row["group"]["layer"] for row in by_layer["rows"]} == {"Takeoff"}
    assert tr.takeoff_summary(path, subject="riprap")["totals"][0]["total"] == pytest.approx(10121.53, abs=0.01)
    assert tr.takeoff_summary(path, group_by="page,kind")["group_by"] == ["page", "kind"]
    assert tr.takeoff_summary(path, layer="review")["rows"] == []
    with pytest.raises(DocumentError, match="Unknown group_by"):
        tr.takeoff_summary(path, group_by=["colour"])


# =============================================================================
# export
# =============================================================================

def _read_csv(p):
    with open(p, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def test_export_csv(synth1, tmp_path):
    path, exp = synth1
    before = open(path, "rb").read()
    out = str(tmp_path / "m.csv")
    r = tr.export_markups(path, out, "csv")
    assert r["rows"] == 13 and r["format"] == "csv" and r["columns"][:8] == tr.EXPORT_COLUMNS[:8]
    assert r["columns"][-3:] == ["Bid Item", "Unit Cost", "Phase"] and r["sheets"] is None
    rows = _read_csv(out)
    assert len(rows) == 13 and list(rows[0]) == r["columns"]
    riprap = next(x for x in rows if x["Subject"] == "Riprap Area" and x["Bid Item"])
    assert (riprap["Area"], riprap["Volume"], riprap["Volume Unit"], riprap["Depth"]) == ("5208.33", "96.45", "cu yd", "6 in")
    assert riprap["Measurement"] == "5,208.33 sf" and riprap["Unit"] == "sf" and riprap["Page"] == "2"
    assert riprap["Page Label"] == "[2] 2 SITE PLAN" and riprap["Layer"] == "Takeoff" and riprap["Color"] == "#FF0000"
    assert riprap["Phase"] == "Phase 1" and riprap["Author"] == S.AUTHOR and riprap["Date"].startswith("2026-09-28")
    line = next(x for x in rows if x["Type"] == "Line")
    assert line["Length"] == "61.94" and line["Unit"] == "ft" and line["Area"] == ""
    counts = [x for x in rows if x["Subject"] == "Inlet Count"]
    assert sorted(x["Count"] for x in counts) == ["", "", "3"], "only the group's first symbol carries N"
    assert sum(1 for x in counts if x["Measurement"]) == 1
    callout = next(x for x in rows if x["Subject"] == "Cloud+" and x["Type"] == "FreeText")
    assert callout["Status"] == "Accepted" and callout["Comments"] == "Verify wall thickness"
    reply = next(x for x in rows if x["Reply To"])
    assert reply["Comments"] == "Checked against the detail" and reply["Author"] == S.REVIEWER
    assert open(path, "rb").read() == before, "the PDF must not change"


def test_export_xlsx(synth1, tmp_path):
    from openpyxl import load_workbook
    path, exp = synth1
    out = str(tmp_path / "m.xlsx")
    r = tr.export_markups(path, out, "xlsx")
    assert r["sheets"] == ["Markups", "Takeoff"]
    wb = load_workbook(out)
    ws = wb["Markups"]
    header = [c.value for c in ws[1]]
    assert header == r["columns"] and ws.freeze_panes == "A2" and ws.auto_filter.ref == ws.dimensions
    assert ws.max_row == 14 and ws[1][0].font.bold
    assert all(ws.column_dimensions[chr(65 + i)].width >= 8 for i in range(5))
    area_col = header.index("Area") + 1
    values = [ws.cell(row=i, column=area_col).value for i in range(2, 15)]
    assert 5208.33 in values and all(isinstance(v, (int, float)) for v in values if v not in (None, ""))
    tk = wb["Takeoff"]
    cells = [[c.value for c in row] for row in tk.iter_rows()]
    assert cells[0] == ["Subject", "Kind", "Markups", "Total", "Unit", "Pages"]
    assert ["Riprap Area", "area", 2, 10121.53, "sf", "2"] in cells
    assert ["Riprap Area", "volume", 1, 96.45, "cu yd", "2"] in cells
    assert any(row[0] == "TOTAL" and row[1] == "area" and row[3] == pytest.approx(exp["totals"]["area_sf"]) for row in cells)


def test_export_json_and_pages_filter(synth1, tmp_path):
    path, _ = synth1
    out = str(tmp_path / "m.json")
    r = tr.export_markups(path, out, "json", pages="3")
    data = json.loads(open(out, encoding="utf-8").read())
    assert r["rows"] == len(data["markups"]) == 5                       # callout, cloud, text, rect + the reply
    assert data["file"] == "synth1.pdf" and data["takeoff"]["rows"] == []
    full = tr.export_markups(path, str(tmp_path / "all.json"), "json")
    assert json.loads(open(tmp_path / "all.json", encoding="utf-8").read())["takeoff"]["totals"]
    assert full["rows"] == 13


def test_export_refuses_bad_targets(synth1, tmp_path):
    path, _ = synth1
    out = str(tmp_path / "m.csv")
    tr.export_markups(path, out, "csv")
    with pytest.raises(DocumentError, match="already exists"):
        tr.export_markups(path, out, "csv")
    tr.export_markups(path, out, "csv", pages="1", overwrite=True)
    assert len(_read_csv(out)) == 0                                     # replaced (page 1 has no markups)
    with pytest.raises(DocumentError, match=r"must end in \.xlsx"):
        tr.export_markups(path, str(tmp_path / "m.csv"), "xlsx")
    with pytest.raises(DocumentError, match=r"never a \.pdf"):
        tr.export_markups(path, path, "csv")
    with pytest.raises(DocumentError, match="folder does not exist"):
        tr.export_markups(path, str(tmp_path / "nope" / "m.csv"), "csv")
    with pytest.raises(DocumentError, match="format"):
        tr.export_markups(path, str(tmp_path / "m.txt"), "txt")
    with pytest.raises(DocumentError, match="Refusing to write"):
        tr.export_markups(path, r"C:\Windows\bb_export_test.csv", "csv")


# =============================================================================
# compare
# =============================================================================

def test_compare_markups(tmp_path):
    a_path = str(tmp_path / "a.pdf")
    exp = S.build_takeoff_pdf(a_path)
    b_path = str(tmp_path / "b.pdf")
    with pymupdf.open(a_path) as doc:
        m = exp["markups"]
        by_subject = {v["subject"]: v for v in m.values() if v["subject"]}
        doc.xref_set_key(by_subject["Text Box"]["xref"], "Contents", "(Changed note)")
        doc.xref_set_key(by_subject["Concrete Slab"]["xref"], "Subj", "(Renamed slab)")
        rect_v = by_subject["Rectangle"]
        page3 = doc[2]
        next(a for a in page3.annots() if a.xref == rect_v["xref"]).set_rect(pymupdf.Rect(360, 410, 510, 510))
        status_x = m[exp["status"][rect_v["id"]]["id"]]["xref"]
        doc.xref_set_key(status_x, "State", "(Accepted)")
        doc.xref_set_key(by_subject["Concrete Slab"]["xref"], "Contents", "(4,000.00 sf)")               # a new Bluebeam value
        removed = by_subject["Length Measurement"]
        page2 = doc[1]
        page2.delete_annot(next(a for a in page2.annots() if a.xref == removed["xref"]))
        new_x = S.add_rectangle(doc, page3, (20, 20, 60, 60), subject="Added Box")
        added_nm = S.nm_of(doc, new_x)
        doc.save(b_path)
    r = tr.compare_markups(a_path, b_path)
    assert r["counts"] == {"a_total": 13, "b_total": 13, "added": 1, "removed": 1, "changed": 3, "unchanged": 9}
    assert [x["id"] for x in r["added"]] == [added_nm] and r["added"][0]["subject"] == "Added Box"
    assert [x["id"] for x in r["removed"]] == [removed["id"]]
    changes = {c["subject"]: c["changes"] for c in r["changed"]}
    assert changes["Text Box"]["contents"] == {"a": "General note", "b": "Changed note"}
    assert changes["Renamed slab"]["subject"] == {"a": "Concrete Slab", "b": "Renamed slab"}
    assert changes["Renamed slab"]["measurement"]["b"] > changes["Renamed slab"]["measurement"]["a"]
    assert changes["Rectangle"]["status"] == {"a": "Rejected", "b": "Accepted"}
    assert changes["Rectangle"]["rect"]["b"] != changes["Rectangle"]["rect"]["a"]
    assert "truncated" not in r
    same = tr.compare_markups(a_path, a_path)
    assert same["counts"]["added"] == same["counts"]["removed"] == same["counts"]["changed"] == 0
    assert same["counts"]["unchanged"] == 13


def test_compare_ignores_sub_point_moves_and_caps_lists(tmp_path):
    doc, a = _pdf(tmp_path, "ca.pdf")
    rects = [doc[0].add_rect_annot(pymupdf.Rect(10, 10 + 5 * i, 20, 14 + 5 * i)) for i in range(60)]
    for i, r in enumerate(rects):
        doc.xref_set_key(r.xref, "NM", S.pdf_str(f"nm-{i}"))
    doc.save(a)
    for i, r in enumerate(rects):
        if i == 0:
            r.set_rect(pymupdf.Rect(10.4, 10.4, 20.4, 14.4))                 # moved 0.4 pt: not a change
        elif i > 1:
            r.set_rect(pymupdf.Rect(30, 10 + 5 * i, 40, 14 + 5 * i))         # moved 20 pt
    b = str(tmp_path / "cb.pdf")
    doc.save(b)
    doc.close()
    r = tr.compare_markups(a, b)
    assert r["counts"]["changed"] == 58 and len(r["changed"]) == tr.LIST_CAP and r["truncated"] is True


def test_compare_markups_without_nm_match_by_position(tmp_path):
    doc, a = _pdf(tmp_path, "pa.pdf")
    r = doc[0].add_rect_annot(pymupdf.Rect(10, 10, 50, 50))
    doc.xref_set_key(r.xref, "NM", "null")
    doc.save(a)
    doc.close()
    doc2, b = _pdf(tmp_path, "pb.pdf", pages=1)
    for rect in ((10, 10, 50, 50), (100, 100, 150, 150)):
        x = doc2[0].add_rect_annot(pymupdf.Rect(*rect))
        doc2.xref_set_key(x.xref, "NM", "null")
    doc2.save(b)
    doc2.close()
    r = tr.compare_markups(a, b)
    assert (r["counts"]["added"], r["counts"]["removed"], r["counts"]["unchanged"]) == (1, 0, 1)


# =============================================================================
# text: search, page text
# =============================================================================

@pytest.fixture(scope="module")
def text_pdf(tmp_path_factory):
    path = str(tmp_path_factory.mktemp("txt") / "text.pdf")
    doc = pymupdf.open()
    p1 = doc.new_page(width=612, height=792)
    p1.insert_text((72, 100), "Sheet A-101 Site Plan", fontsize=12)
    p1.insert_text((72, 130), "Manhole MH-4 at station 12+50", fontsize=12)
    p2 = doc.new_page(width=612, height=792)
    p2.insert_text((72, 100), "sheet a-101 grading", fontsize=12)
    p2.insert_text((72, 130), "Contact SITE PLAN office; MH-12 near the gate", fontsize=12)
    doc.new_page(width=612, height=792)                                   # page 3 has no text
    doc.save(path)
    doc.close()
    return path


def test_search_text_plain(text_pdf):
    r = tr.search_text(text_pdf, "a-101")
    assert r["total_hits"] == 2 and [h["page"] for h in r["hits"]] == [1, 2]
    h = r["hits"][0]
    assert len(h["rect"]) == 4 and 70 < h["rect"][0] < 130 and 85 < h["rect"][1] < 105
    assert "Sheet A-101 Site Plan" in h["context"]
    assert tr.search_text(text_pdf, "site plan")["total_hits"] == 2
    assert tr.search_text(text_pdf, "site plan", pages="2")["total_hits"] == 1
    assert [h["page"] for h in tr.search_text(text_pdf, "a-101", case_sensitive=True)["hits"]] == [2]
    assert [h["page"] for h in tr.search_text(text_pdf, "A-101", case_sensitive=True)["hits"]] == [1]
    assert tr.search_text(text_pdf, "a-102", case_sensitive=True)["total_hits"] == 0
    assert tr.search_text(text_pdf, "Site Plan", case_sensitive=True)["total_hits"] == 1


def test_search_text_regex_and_limits(text_pdf):
    r = tr.search_text(text_pdf, r"MH-\d+", regex=True)
    assert [h["context"].split(" ")[0] for h in r["hits"]] == ["Manhole", "Contact"] and r["total_hits"] == 2
    assert tr.search_text(text_pdf, r"^mh-4$", regex=True, case_sensitive=True)["total_hits"] == 0
    assert tr.search_text(text_pdf, r"^mh-4$", regex=True)["total_hits"] == 1
    capped = tr.search_text(text_pdf, "a", max_hits=3)
    assert capped["returned"] == 3 and capped["total_hits"] > 3 and capped["truncated"] is True
    with pytest.raises(DocumentError, match="Bad regex"):
        tr.search_text(text_pdf, "(", regex=True)
    with pytest.raises(DocumentError, match="empty"):
        tr.search_text(text_pdf, "")


def test_search_text_no_text_layer_note(text_pdf, tmp_path):
    r = tr.search_text(text_pdf, "zzz")
    assert r["total_hits"] == 0 and r["hits"] == [] and "1 of 3" in r["note"]
    doc, path = _pdf(tmp_path, "blank.pdf", pages=2)
    _save(doc, path)
    assert "2 of 2" in tr.search_text(path, "anything")["note"]


def test_get_page_text_modes(text_pdf):
    t = tr.get_page_text(text_pdf, 1)
    assert "Sheet A-101 Site Plan" in t["text"] and "MH-4" in t["text"] and t["mode"] == "text"
    blocks = tr.get_page_text(text_pdf, 1, mode="blocks")["blocks"]
    assert blocks and set(blocks[0]) == {"rect", "text"} and "A-101" in blocks[0]["text"]
    words = tr.get_page_text(text_pdf, 1, mode="words")["words"]
    assert len(words[0]) == 5 and words[0][4] == "Sheet" and [w[4] for w in words][:3] == ["Sheet", "A-101", "Site"]
    clipped = tr.get_page_text(text_pdf, 1, clip=[0, 115, 612, 140])["text"]
    assert "Manhole" in clipped and "Sheet" not in clipped
    empty = tr.get_page_text(text_pdf, 3)
    assert empty["text"] == "" and "scanned" in empty["note"]
    for bad in ({"mode": "html"}, {"clip": [1, 2, 3]}):
        with pytest.raises(DocumentError):
            tr.get_page_text(text_pdf, 1, **bad)
    with pytest.raises(DocumentError, match="out of range"):
        tr.get_page_text(text_pdf, 4)


def test_get_page_text_truncation(tmp_path):
    doc, path = _pdf(tmp_path, "long.pdf")
    for i in range(80):
        doc[0].insert_text((20, 20 + i * 9), f"Line number {i:03d} of a long note on this sheet", fontsize=7)
    _save(doc, path)
    t = tr.get_page_text(path, 1, max_chars=300)
    assert t["truncated"] is True and len(t["text"]) == 300
    w = tr.get_page_text(path, 1, mode="words", max_chars=300)
    assert w["truncated"] is True and 0 < len(w["words"]) < 100
    b = tr.get_page_text(path, 1, mode="blocks", max_chars=300)
    assert "truncated" in b or len(b["blocks"]) == 1
    assert "truncated" not in tr.get_page_text(path, 1, max_chars=200000)


# =============================================================================
# sheet index and layers tool
# =============================================================================

def test_sheet_index(synth1, tmp_path):
    src, exp = synth1
    path = str(tmp_path / "idx.pdf")
    with pymupdf.open(src) as doc:
        doc[1].insert_text((470, 700), "C-101", fontsize=30)
        doc[1].insert_text((470, 740), "SITE PLAN", fontsize=14)
        doc[1].insert_text((470, 760), "rev", fontsize=6)
        doc[0].insert_text((50, 50), "COVER", fontsize=40)                # not in the bottom-right zone
        doc.set_toc([[1, "Cover", 1], [1, "Site Plan", 2], [2, "Grading", 2]])
        doc[2].set_rotation(90)
        doc.save(path)
    r = tr.sheet_index(path)
    assert r["page_count"] == 3 and r["markups"] == 12
    p1, p2, p3 = r["pages"]
    assert p1["number"] == 1 and p1["label"] == "[1] 1 COVER SHEET" and p1["bookmark"] == "Cover"
    assert p1["size_in"] == [8.5, 11.0] and p1["orientation"] == "portrait" and p1["markups"] == 0
    assert "title_block" not in p1 and "scales" not in p1
    assert p2["bookmark"] == "Site Plan" and p2["scales"] == ["1 in = 30 ft' in\""] and p2["markups"] == 8
    assert p2["title_block"] == ["C-101", "SITE PLAN", "rev"]
    assert p3["rotation"] == 90 and p3["orientation"] == "landscape" and p3["size_in"] == [11.0, 8.5]
    assert p3["markups"] == 4 and "bookmark" not in p3


def test_sheet_index_title_block_on_rotated_page(tmp_path):
    doc, path = _pdf(tmp_path, "rotidx.pdf")
    page = doc[0]
    # displayed (rotated 90) bottom-right corner corresponds to unrotated top-right region
    page.set_rotation(90)
    zone = pymupdf.Rect(page.rect.width * 0.8, page.rect.height * 0.8, page.rect.width - 5, page.rect.height - 5)
    spot = (zone.tl * page.derotation_matrix)
    page.insert_text((spot.x, spot.y), "X-9", fontsize=20, rotate=270)
    _save(doc, path)
    (p,) = tr.sheet_index(path)["pages"]
    assert p["rotation"] == 90 and p.get("title_block") == ["X-9"]


def test_sheet_index_survives_damaged_outline(sample_pdf_3p, monkeypatch):
    monkeypatch.setattr(pymupdf.Document, "get_toc", lambda self, *a, **k: (_ for _ in ()).throw(RuntimeError("bad outline")))
    r = tr.sheet_index(sample_pdf_3p)
    assert r["page_count"] == 3 and all("bookmark" not in p for p in r["pages"])


def test_list_layers_tool(synth1, sample_pdf):
    r = tr.list_layers(synth1[0])
    assert r["count"] == 2 and [l["name"] for l in r["layers"]] == ["Takeoff", "Review"] and "note" not in r
    none = tr.list_layers(sample_pdf)
    assert none["count"] == 0 and "no layers" in none["note"]


# =============================================================================
# MCP end to end, empty PDFs, speed
# =============================================================================

ALL_TOOLS = {"bb_list_markups", "bb_get_markup", "bb_takeoff_summary", "bb_export_markups", "bb_search_text",
             "bb_get_page_text", "bb_sheet_index", "bb_render_page", "bb_render_markup", "bb_compare_markups",
             "bb_list_layers"}


def test_registration_and_annotations():
    tools = {t.name: t for t in asyncio.run(_mcp().list_tools())}
    assert set(tools) == ALL_TOOLS
    for name, t in tools.items():
        assert t.annotations.openWorldHint is False and t.description and len(t.description) > 60
        if name == "bb_export_markups":
            assert t.annotations.readOnlyHint is False and t.annotations.destructiveHint is False
        else:
            assert t.annotations.readOnlyHint is True
        assert t.outputSchema is None, "compact JSON text only, no duplicated structured payload"
    assert set(tools["bb_list_markups"].inputSchema["properties"]) == {
        "path", "pages", "types", "subject", "author", "layer", "include_replies", "detail", "limit"}


def test_tools_end_to_end_via_mcp(synth1, tmp_path):
    path, exp = synth1
    r = _call("bb_list_markups", path=path, pages="2", types=["area"], detail="summary")
    assert r["total"] == 3 and all(m["measurement"]["kind"] == "area" for m in r["markups"])
    assert _call("bb_list_markups", path=path, pages=3, include_replies=True)["total"] == 5
    t = _call("bb_takeoff_summary", path=path, group_by=["subject", "layer"])
    assert t["group_by"] == ["subject", "layer"] and t["warnings"] == []
    out = str(tmp_path / "e2e.xlsx")
    e = _call("bb_export_markups", path=path, output_path=out, format="xlsx")
    assert e["rows"] == 13 and (tmp_path / "e2e.xlsx").exists()
    nm = next(k for k, v in exp["markups"].items() if v["kind"] == "callout")
    assert _call("bb_get_markup", path=path, markup_id=nm)["markup"]["subject"] == "Cloud+"
    assert _call("bb_list_layers", path=path)["count"] == 2
    assert _call("bb_sheet_index", path=path)["page_count"] == 3
    assert _call("bb_compare_markups", path_a=path, path_b=path)["counts"]["changed"] == 0
    assert _call("bb_search_text", path=path, query="zzz")["total_hits"] == 0
    assert "Verify wall thickness" in _call("bb_get_page_text", path=path, page=3)["text"]
    with pytest.raises(Exception, match="Markup not found"):
        _call("bb_get_markup", path=path, markup_id="missing")
    with pytest.raises(Exception, match="already exists"):
        _call("bb_export_markups", path=path, output_path=out, format="xlsx")


def test_every_tool_on_a_pdf_without_annotations(sample_pdf, tmp_path):
    assert _call("bb_list_markups", path=sample_pdf)["total"] == 0
    assert _call("bb_takeoff_summary", path=sample_pdf) == {
        "file": "sample.pdf", "group_by": ["subject"], "rows": [], "totals": [], "scales": [], "warnings": [],
        "measured": 0, "skipped": 0}
    assert _call("bb_list_layers", path=sample_pdf)["count"] == 0
    idx = _call("bb_sheet_index", path=sample_pdf)
    assert idx["page_count"] == 1 and idx["pages"][0]["markups"] == 0
    assert _call("bb_search_text", path=sample_pdf, query="Sample")["total_hits"] == 1
    assert "Sample page 1" in _call("bb_get_page_text", path=sample_pdf, page=1)["text"]
    assert _call("bb_compare_markups", path_a=sample_pdf, path_b=sample_pdf)["counts"]["a_total"] == 0
    for fmt in ("csv", "xlsx", "json"):
        r = _call("bb_export_markups", path=sample_pdf, output_path=str(tmp_path / f"none.{fmt}"), format=fmt)
        assert r["rows"] == 0
    assert _read_csv(str(tmp_path / "none.csv")) == []
    with pytest.raises(Exception, match="Markup not found"):
        _call("bb_get_markup", path=sample_pdf, markup_id="x")
    out = asyncio.run(_mcp().call_tool("bb_render_page", {"path": sample_pdf, "page": 1}))
    assert (out[0] if isinstance(out, tuple) else out)[0].type == "image"


def test_bad_inputs_raise_tool_errors(tmp_path):
    with pytest.raises(Exception, match="File not found"):
        _call("bb_list_markups", path=str(tmp_path / "missing.pdf"))
    (tmp_path / "not.pdf").write_bytes(b"this is not a pdf")
    with pytest.raises(Exception, match="Could not open PDF"):
        _call("bb_list_markups", path=str(tmp_path / "not.pdf"))
    with pytest.raises(Exception, match="Not a .pdf"):
        _call("bb_list_layers", path=str(tmp_path / "x.txt"))


def test_fifty_page_pdf_is_fast(tmp_path):
    path = str(tmp_path / "fifty.pdf")
    doc = pymupdf.open()
    for i in range(50):
        page = doc.new_page(width=2448, height=1584)
        page.insert_text((100, 100), f"Sheet C-{i + 1:03d} grading plan", fontsize=14)
        page.insert_text((2100, 1450), f"C-{i + 1:03d}", fontsize=40)
        S.add_area(doc, page, [(200, 200), (600, 200), (600, 500), (200, 500)], subject="Grading Area")
        page.add_rect_annot(pymupdf.Rect(700, 200, 800, 300))
    doc.save(path)
    doc.close()
    t0 = time.time()
    lst = tr.list_markups(path, limit=2000)
    assert lst["total"] == 100
    ts = tr.takeoff_summary(path)
    assert ts["totals"][0]["total"] == pytest.approx(50 * 20833.35, rel=1e-4)
    idx = tr.sheet_index(path)
    assert idx["page_count"] == 50 and idx["pages"][49]["title_block"][0] == "C-050"
    assert tr.search_text(path, "C-0")["total_hits"] == 100          # 2 per sheet: the note and the title block
    assert tr.search_text(path, "grading plan")["total_hits"] == 50
    tr.export_markups(path, str(tmp_path / "fifty.xlsx"), "xlsx")
    tr.list_layers(path)
    assert time.time() - t0 < 60


def test_blank_viewport_scale_text_is_dropped(tmp_path):
    import pymupdf as _pm
    from bluebeam import markups_read as _mr
    doc = _pm.open()
    page = doc.new_page(width=612, height=792)
    m = doc.get_new_xref()
    doc.update_object(m, "<</Type/Measure/Subtype/RL/R( )/X[<</Type/NumberFormat/U(ft)/C .4166667/D 100>>]"
                         "/D[<</Type/NumberFormat/U(ft)/C 1/D 100>>]/A[<</Type/NumberFormat/U(sf)/C 1/D 100>>]>>")
    doc.xref_set_key(page.xref, "VP", f"[<</Type/Viewport/BBox[0 0 612 792]/Measure {m} 0 R>>]")
    scales = _mr.page_scales(doc, 0)
    assert scales and "scale" not in scales[0] and scales[0]["units_per_point"] == 0.4166667
    from bluebeam.tools_read import _scale_from_factor
    assert _scale_from_factor(scales[0]) == "1 in = 30 ft"
