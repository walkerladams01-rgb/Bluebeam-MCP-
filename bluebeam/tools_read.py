"""Read-only MCP tools: Bluebeam-aware markup reading, takeoff quantities, text search, page rendering.

Logic lives in plain functions at the top (each returns JSON-ready data and is unit-testable);
register_read_tools() wraps them as thin FastMCP tools that return compact JSON text (images
are returned as image content plus a small JSON note with pixel size and dpi).
"""
from __future__ import annotations

import csv
import json
import re
from collections import Counter
from pathlib import Path
from typing import Any, Optional, Union

import pymupdf
from mcp.server.fastmcp import FastMCP, Image
from mcp.types import ToolAnnotations

from . import core, markups_read as mr, render, takeoff

Pages = Union[int, str, list[int], None]
MAX_LIMIT = 2000
LIST_CAP = 50                       # entries in the added/removed/changed lists of a comparison


def _j(obj: Any) -> str:
    return json.dumps(obj, ensure_ascii=False, separators=(",", ":"), default=str)


def _has(hay: Optional[str], needle: Optional[str]) -> bool:
    return not needle or needle.lower() in (hay or "").lower()


def _list(v: Union[list[str], str, None]) -> list[str]:
    if v is None:
        return []
    return [p.strip() for p in v.split(",") if p.strip()] if isinstance(v, str) else [str(p).strip() for p in v]


def _name(path: str) -> str:
    return Path(path).name


def _detail(detail: str) -> str:
    if detail not in ("summary", "full"):
        raise core.DocumentError("detail must be 'summary' or 'full'")
    return detail


# =============================================================================
# markups
# =============================================================================

def _filter(mks: list[mr.Markup], types: list[str], subject: Optional[str], author: Optional[str],
            layer: Optional[str]) -> list[mr.Markup]:
    wanted = {t.lower() for t in types}

    def keep(m: mr.Markup) -> bool:
        if wanted:
            names = {m.pdf_type.lower(), (m.intent or "").lower(), m.kind}
            if m.measurement:
                names.add("measurement")
            if not wanted & names:
                return False
        return _has(m.subject, subject) and _has(m.author, author) and _has(m.layer, layer)

    return [m for m in mks if keep(m)]


def list_markups(path: str, pages: Pages = None, types: Union[list[str], str, None] = None,
                 subject: Optional[str] = None, author: Optional[str] = None, layer: Optional[str] = None,
                 include_replies: bool = False, detail: str = "summary", limit: int = 300) -> dict:
    _detail(detail)
    with core.open_readonly(path) as doc:
        idx = core.parse_page_range(pages, doc.page_count)
        mks = mr.parse_markups(doc, idx, include_replies)
    mks = _filter(mks, _list(types), subject, author, layer)
    limit = max(1, min(int(limit), MAX_LIMIT))
    out: dict[str, Any] = {"file": _name(path), "total": len(mks), "returned": min(len(mks), limit)}
    if len(mks) > limit:
        out["truncated"] = True
        out["note"] = f"showing the first {limit}; narrow with pages/types/subject/layer or raise limit (max {MAX_LIMIT})"
    out["markups"] = [m.to_dict(detail) for m in mks[:limit]]
    return out


def get_markup(path: str, markup_id: str) -> dict:
    with core.open_readonly(path) as doc:
        page, annot = core.find_annot(doc, markup_id)
        mks = mr.parse_markups(doc, [page.number], include_replies=True)
        mk = next((m for m in mks if m.xref == annot.xref), None)
        if mk is None:
            raise core.DocumentError(f"'{markup_id}' is a review-status entry, not a markup; its state is "
                                     "shown as 'status' on the markup it belongs to")
        return {"file": _name(path), "markup": mk.to_dict("full"), "raw_keys": mr.raw_keys(doc, annot.xref),
                "group_children": [m.id for m in mks if m.group_parent_id == mk.id],
                "replies": [{"id": m.id, "author": m.author, "date": m.modified or m.created,
                             "contents": m.contents} for m in mks if m.reply_to_id == mk.id]}


def takeoff_summary(path: str, group_by: Union[list[str], str, None] = None, pages: Pages = None,
                    layer: Optional[str] = None, subject: Optional[str] = None) -> dict:
    with core.open_readonly(path) as doc:
        idx = core.parse_page_range(pages, doc.page_count)
        mks = mr.parse_markups(doc, idx)
    mks = _filter(mks, [], subject, None, layer)
    out = takeoff.summarize(mks, _list(group_by) or ["subject"])
    return {"file": _name(path), **out}


# --- export ------------------------------------------------------------------

EXPORT_COLUMNS = ["Subject", "Page", "Page Label", "Author", "Date", "Status", "Layer", "Comments",
                  "Measurement", "Unit", "Length", "Area", "Volume", "Volume Unit", "Count", "Depth",
                  "Label", "Color", "Fill Color", "Type", "Id", "Reply To"]


def _export_row(m: mr.Markup) -> dict[str, Any]:
    me = m.measurement
    counted = me if me and not me.get("member_of") else None       # count symbols: the group's row holds N
    r = lambda v: round(v, 2) if isinstance(v, float) else v        # noqa: E731
    row: dict[str, Any] = {
        "Subject": m.subject, "Page": m.page, "Page Label": m.page_label, "Author": m.author,
        "Date": m.modified or m.created or "", "Status": m.status["state"] if m.status else "",
        "Layer": m.layer or "", "Comments": (m.rich_text_plain or "") if me else (m.contents or m.rich_text_plain),
        "Measurement": (me.get("formatted") or "") if counted else "", "Unit": (me.get("unit") or "") if counted else "",
        "Length": r(me["value"]) if counted and me["kind"] in ("length", "perimeter") else "",
        "Area": r(me["value"]) if counted and me["kind"] == "area" else "",
        "Volume": r(me["volume"]) if counted and me.get("volume") is not None else "",
        "Volume Unit": (me.get("volume_unit") or "") if counted and me.get("volume") is not None else "",
        "Count": me["count"] if counted and me["kind"] == "count" else "",
        "Depth": f"{me['depth']:g} {me.get('depth_unit') or ''}".strip() if counted and me.get("depth") is not None else "",
        "Label": m.label, "Color": m.color or "", "Fill Color": m.fill_color or "", "Type": m.pdf_type, "Id": m.id, "Reply To": m.reply_to_id or ""}
    row.update(m.custom_columns)
    return row


def _swatch(cells) -> None:
    """Shade cells holding '#RRGGBB' in that colour (white text on dark colours)."""
    from openpyxl.styles import Font, PatternFill
    for cell in cells:
        v = str(cell.value or "")
        if len(v) == 7 and v.startswith("#"):
            try:
                r, g, b = (int(v[i:i + 2], 16) for i in (1, 3, 5))
            except ValueError:
                continue
            cell.fill = PatternFill("solid", fgColor=v[1:].upper())
            if 0.299 * r + 0.587 * g + 0.114 * b < 128:
                cell.font = Font(color="FFFFFF")


def _write_xlsx(out: Path, columns: list[str], rows: list[dict], summary: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Markups"
    ws.append(columns)
    for row in rows:
        ws.append([row.get(c, "") for c in columns])
    head = PatternFill("solid", fgColor="D9E1F2")
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True), head
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(columns, 1):
        longest = max([len(str(col))] + [len(str(r.get(col, ""))) for r in rows[:500]])
        ws.column_dimensions[get_column_letter(i)].width = max(8, min(longest + 2, 60))
    for cell in ws[get_column_letter(columns.index("Comments") + 1)][1:]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for name in ("Color", "Fill Color"):
        _swatch(ws[get_column_letter(columns.index(name) + 1)][1:])
    for name in ("Length", "Area", "Volume"):
        for cell in ws[get_column_letter(columns.index(name) + 1)][1:]:
            cell.number_format = "#,##0.00"

    tk = wb.create_sheet("Takeoff")
    group_cols = [g.replace("_", " ").title() for g in summary["group_by"]]
    tk.append(group_cols + ["Kind", "Markups", "Total", "Unit", "Pages"])
    for r in summary["rows"]:
        tk.append([r["group"][g] for g in summary["group_by"]] + [r["kind"], r["markups"], r["total"], r["unit"],
                                                                 ", ".join(map(str, r["pages"]))])
    if summary["totals"]:
        tk.append([])
        for t in summary["totals"]:
            tk.append(["TOTAL"] + [""] * (len(group_cols) - 1) + [t["kind"], t["markups"], t["total"], t["unit"], ""])
    for w in summary["warnings"]:
        tk.append([])
        tk.append(["WARNING: " + w])
    for cell in tk[1]:
        cell.font, cell.fill = Font(bold=True), head
    tk.freeze_panes = "A2"
    for i, g in enumerate(summary["group_by"], 1):
        if g in ("color", "fill_color"):
            _swatch(tk[get_column_letter(i)][1:])
    for cell in tk[get_column_letter(len(group_cols) + 3)][1:]:
        cell.number_format = "#,##0.00"
    for i in range(1, len(group_cols) + 6):
        tk.column_dimensions[get_column_letter(i)].width = 28 if i <= len(group_cols) else 14
    wb.save(str(out))


def export_markups(path: str, output_path: str, format: str = "csv", pages: Pages = None,
                   overwrite: bool = False, group_by: Union[list[str], str, None] = None) -> dict:
    fmt = (format or "").lower()
    if fmt not in ("csv", "xlsx", "json"):
        raise core.DocumentError("format must be 'csv', 'xlsx' or 'json'")
    if not output_path or not output_path.strip():
        raise core.DocumentError("No output_path given")
    out = Path(output_path.strip().strip('"')).expanduser().absolute()
    if out.suffix.lower() != f".{fmt}":
        raise core.DocumentError(f"output_path must end in .{fmt} (never a .pdf): {out.name}")
    core.check_writable_target(out)
    if not out.parent.is_dir():
        raise core.DocumentError(f"Output folder does not exist: {out.parent}")
    if out.exists() and not overwrite:
        raise core.DocumentError(f"Output already exists: {out} (pass overwrite=True to replace it)")
    with core.open_readonly(path) as doc:
        idx = core.parse_page_range(pages, doc.page_count)
        mks = mr.parse_markups(doc, idx, include_replies=True)
    summary = takeoff.summarize(mks, _list(group_by) or ["subject"])
    custom = list(dict.fromkeys(name for m in mks for name in m.custom_columns))
    columns = EXPORT_COLUMNS + custom
    rows = [_export_row(m) for m in mks]
    if fmt == "csv":
        with open(out, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
            w.writeheader()
            w.writerows(rows)
    elif fmt == "xlsx":
        _write_xlsx(out, columns, rows, summary)
    else:
        out.write_text(json.dumps({"file": _name(path), "markups": [m.to_dict("full") for m in mks],
                                   "takeoff": summary}, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    return {"output_path": str(out), "format": fmt, "rows": len(rows), "columns": columns,
            "sheets": ["Markups", "Takeoff"] if fmt == "xlsx" else None}


# --- compare -----------------------------------------------------------------

def _key(m: mr.Markup) -> str:
    """Identity across two files: the /NM GUID, else page + type + position."""
    if not m.id.startswith("xref:"):
        return m.id
    return f"pos:{m.page}:{m.pdf_type}:" + ",".join(f"{v:.0f}" for v in (m.rect or []))


def _diff(a: mr.Markup, b: mr.Markup) -> dict[str, dict]:
    changes: dict[str, dict] = {}
    for name, va, vb in (("contents", a.contents, b.contents), ("subject", a.subject, b.subject),
                         ("author", a.author, b.author), ("layer", a.layer or "", b.layer or ""),
                         ("status", a.status["state"] if a.status else "", b.status["state"] if b.status else ""),
                         ("measurement", _value(a), _value(b))):
        if va != vb:
            changes[name] = {"a": va, "b": vb}
    if a.rect and b.rect and max(abs(x - y) for x, y in zip(a.rect, b.rect)) > 1:
        changes["rect"] = {"a": [round(v, 1) for v in a.rect], "b": [round(v, 1) for v in b.rect]}
    return changes


def _value(m: mr.Markup) -> Optional[float]:
    v = m.measurement.get("value") if m.measurement else None
    return round(v, 2) if isinstance(v, float) else v


def compare_markups(path_a: str, path_b: str) -> dict:
    with core.open_readonly(path_a) as da, core.open_readonly(path_b) as db:
        a = {_key(m): m for m in mr.parse_markups(da, include_replies=True)}
        b = {_key(m): m for m in mr.parse_markups(db, include_replies=True)}
    added = [m for k, m in b.items() if k not in a]
    removed = [m for k, m in a.items() if k not in b]
    changed = []
    for k in a.keys() & b.keys():
        diff = _diff(a[k], b[k])
        if diff:
            changed.append({"id": b[k].id, "page": b[k].page, "subject": b[k].subject or b[k].pdf_type,
                            "changes": diff})
    changed.sort(key=lambda c: (c["page"], c["id"]))
    brief = lambda m: {k: v for k, v in m.to_dict().items()                        # noqa: E731
                       if k in ("id", "page", "type", "subject", "author", "contents", "layer", "measurement")}
    capped = lambda items: items[:LIST_CAP]                                        # noqa: E731
    return {"a": _name(path_a), "b": _name(path_b),
            "counts": {"a_total": len(a), "b_total": len(b), "added": len(added), "removed": len(removed),
                       "changed": len(changed), "unchanged": len(a.keys() & b.keys()) - len(changed)},
            "added": [brief(m) for m in capped(added)], "removed": [brief(m) for m in capped(removed)],
            "changed": capped(changed),
            **({"truncated": True} if any(len(x) > LIST_CAP for x in (added, removed, changed)) else {})}


def list_layers(path: str) -> dict:
    with core.open_readonly(path) as doc:
        layers = mr.list_layers(doc)
    out: dict[str, Any] = {"file": _name(path), "count": len(layers), "layers": layers}
    if not layers:
        out["note"] = "this PDF has no layers (optional content groups)"
    return out


# =============================================================================
# text
# =============================================================================

def _squash(s: str) -> str:
    return " ".join(s.split())


def _rect(r: Any) -> list[float]:
    return [round(float(v), 1) for v in (r[0], r[1], r[2], r[3])]


def _context(words: list, hit: pymupdf.Rect, width: int = 160) -> str:
    """The text line a hit sits on (its words in reading order), trimmed around the hit to ~width chars."""
    idx = next((i for i, w in enumerate(words) if w[0] < hit.x1 and w[2] > hit.x0 and w[1] < hit.y1 and w[3] > hit.y0),
               None)
    if idx is None:
        return ""
    block, line = words[idx][5], words[idx][6]
    same = [w for w in words if w[5] == block and w[6] == line]
    text = " ".join(w[4] for w in same)
    if len(text) > width:
        pos = next(i for i, w in enumerate(same) if w is words[idx])
        text = " ".join(w[4] for w in same[max(0, pos - 8):pos + 9])
    return text[:width]


def search_text(path: str, query: str, pages: Pages = None, case_sensitive: bool = False,
                max_hits: int = 200, regex: bool = False) -> dict:
    if not query:
        raise core.DocumentError("query is empty")
    max_hits = max(1, min(int(max_hits), 1000))
    rx = None
    if regex:
        try:
            rx = re.compile(query, 0 if case_sensitive else re.IGNORECASE)
        except re.error as e:
            raise core.DocumentError(f"Bad regex '{query}': {e}") from None
    hits: list[dict] = []
    total = 0
    with core.open_readonly(path) as doc:
        idx = core.parse_page_range(pages, doc.page_count)
        for pno in idx:
            page = doc[pno]
            tp = page.get_textpage()
            words = page.get_text("words", textpage=tp)
            if rx:
                rects = [w[:4] for w in words if rx.search(w[4])]
            else:
                rects = page.search_for(query, textpage=tp)
                if case_sensitive:
                    rects = [r for r in rects if query in _squash(page.get_textbox(r, textpage=tp))]
            total += len(rects)
            label = core.decode_label(page.get_label())
            for r in rects:
                if len(hits) >= max_hits:
                    break
                r = pymupdf.Rect(r)
                hit = {"page": pno + 1, "rect": _rect(r), "context": _context(words, r)}
                if label:
                    hit["page_label"] = label
                hits.append(hit)
        no_text = 0 if total else sum(1 for p in idx if not doc[p].get_text("text").strip())
    out: dict[str, Any] = {"file": _name(path), "query": query, "total_hits": total, "returned": len(hits),
                           "hits": hits}
    if total > len(hits):
        out["truncated"] = True
    if not total and no_text:
        out["note"] = f"{no_text} of {len(idx)} searched page(s) have no text layer (scanned/raster); search cannot read them"
    return out


def get_page_text(path: str, page: int, mode: str = "text", clip: Optional[list[float]] = None,
                  max_chars: int = 20000) -> dict:
    if mode not in ("text", "blocks", "words"):
        raise core.DocumentError("mode must be 'text', 'blocks' or 'words'")
    max_chars = max(200, min(int(max_chars), 200000))
    if clip is not None and len(clip) != 4:
        raise core.DocumentError("clip must be [x0, y0, x1, y1] in PDF points")
    with core.open_readonly(path) as doc:
        pg = doc[core.page_index(doc, page)]
        area = pymupdf.Rect(*clip) if clip else None
        out: dict[str, Any] = {"file": _name(path), "page": page, "mode": mode}
        label = core.decode_label(pg.get_label())
        if label:
            out["page_label"] = label
        if mode == "text":
            text = pg.get_text("text", clip=area)
            out["truncated"] = len(text) > max_chars or None
            out["text"] = text[:max_chars]
            count = len(text)
        else:
            raw = pg.get_text(mode, clip=area)
            items, used = [], 0
            for t in raw:
                if mode == "blocks" and t[6] != 0:          # skip image blocks
                    continue
                item = {"rect": _rect(t), "text": t[4].strip()} if mode == "blocks" else [*_rect(t), t[4]]
                used += len(t[4]) + 12
                if used > max_chars:
                    out["truncated"] = True
                    break
                items.append(item)
            out[mode] = items
            count = len(raw)
        if not count:
            out["note"] = "no text found here (scanned/raster page, or an empty clip); use bb_render_page to look at it"
    return {k: v for k, v in out.items() if v is not None}


# =============================================================================
# sheets
# =============================================================================

def _title_block(page: pymupdf.Page) -> list[str]:
    """Largest-font texts in the bottom-right quarter of the displayed page (sheet number/title guess)."""
    r = page.rect
    zone = pymupdf.Rect(r.x0 + r.width * 0.75, r.y0 + r.height * 0.75, r.x1, r.y1)
    if page.rotation:
        zone = pymupdf.Rect(zone * page.derotation_matrix).normalize()      # text space is unrotated
    spans = []
    for block in page.get_text("dict", clip=zone).get("blocks", []):
        for line in block.get("lines", []):
            for sp in line.get("spans", []):
                text = sp["text"].strip()
                if text:
                    spans.append((sp["size"], text))
    seen: list[str] = []
    for _, text in sorted(spans, key=lambda s: -s[0]):
        if text not in seen:
            seen.append(text)
    return seen[:3]


def sheet_index(path: str) -> dict:
    with core.open_readonly(path) as doc:
        mks = mr.parse_markups(doc)
        per_page = Counter(m.page for m in mks)
        titles: dict[int, str] = {}
        try:
            for _lvl, title, pno, *_ in doc.get_toc():
                if pno > 0:
                    titles.setdefault(pno, title)
        except Exception as e:                                # a damaged outline must not fail the index
            core.log.debug("could not read outline: %s", e)
        pages = []
        for pno in range(doc.page_count):
            page = doc[pno]
            r = page.rect
            scales = [x for x in dict.fromkeys(s.get("scale") or _scale_from_factor(s)
                                               for s in mr.page_scales(doc, pno, mks)) if x]
            entry: dict[str, Any] = {
                "number": pno + 1, "label": core.decode_label(page.get_label()), "bookmark": titles.get(pno + 1, ""),
                "size_in": [round(r.width / 72, 2), round(r.height / 72, 2)],
                "orientation": "landscape" if r.width > r.height else "portrait",
                "rotation": page.rotation, "scales": scales, "title_block": _title_block(page),
                "markups": per_page.get(pno + 1, 0)}
            pages.append({k: v for k, v in entry.items() if v not in ("", [], 0) or k in ("number", "markups")})
        return {"file": _name(path), "page_count": doc.page_count, "markups": len(mks), "pages": pages}


# =============================================================================
# images
# =============================================================================

def _png_result(r: render.Rendered, **meta: Any) -> list:
    return [Image(data=r.png, format="png"), _j({**meta, **r.meta(), "px_per_pt": round(r.dpi / 72, 4)})]


def render_page_tool(path: str, page: int, max_px: int = 1600, clip: Optional[list[float]] = None,
                     annots: bool = True) -> list:
    with core.open_readonly(path) as doc:
        idx = core.page_index(doc, page)
        r = render.render_page_info(doc, idx, max_px, clip, annots)
    return _png_result(r, page=page)


def render_markup_tool(path: str, markup_id: str, margin: float = 36, dpi: float = 150) -> list:
    with core.open_readonly(path) as doc:
        page, annot = core.find_annot(doc, markup_id)
        r = render.render_markup_info(doc, page, annot, margin, dpi)
        return _png_result(r, page=page.number + 1, markup_id=core.markup_id(doc, annot))


# =============================================================================
# registration
# =============================================================================

def _scale_from_factor(s: dict) -> Optional[str]:
    """'1 in = 30 ft' from units_per_point when the PDF carries no scale text; None when the
    units are unnamed too (e.g. AutoCAD viewport factors), which says nothing useful."""
    upp, unit = s.get("units_per_point"), s.get("unit")
    return f"1 in = {upp * 72:.4g} {unit}" if upp and unit else None


def register_read_tools(mcp: FastMCP) -> None:
    """Register the read-only Bluebeam tools on a FastMCP server."""
    ro = ToolAnnotations(readOnlyHint=True, openWorldHint=False)
    export_ann = ToolAnnotations(readOnlyHint=False, destructiveHint=False, openWorldHint=False)

    @mcp.tool(annotations=ro, structured_output=False)
    def bb_list_markups(path: str, pages: Pages = None, types: Union[list[str], str, None] = None,
                        subject: Optional[str] = None, author: Optional[str] = None, layer: Optional[str] = None,
                        include_replies: bool = False, detail: str = "summary", limit: int = 300) -> str:
        """List the markups in a PDF, understanding Bluebeam's structures (read-only).

        pages: 1-based, e.g. 3, "1,3-5", "4-" or omit for all. types: PDF types (Polygon, FreeText, ...)
        or measurement kinds (area, length, perimeter, count, angle, volume, measurement).
        subject/author/layer: case-insensitive substring filters. include_replies also lists comment
        replies (review statuses are never listed; they appear as 'status' on their markup).
        detail "summary" is compact; "full" adds colors, rect, vertices, dates, scale details.
        Coordinates are PDF points, origin top-left (px = pt x dpi / 72). Measurement markups carry
        {kind, value, unit, formatted (Bluebeam's own text)}. A Bluebeam count is one group: each symbol
        repeats the total N and has member_of set, so do not sum members (bb_takeoff_summary does it right).
        Returns {file, total, returned, truncated?, markups[]}; limit max 2000."""
        return _j(list_markups(path, pages, types, subject, author, layer, include_replies, detail, limit))

    @mcp.tool(annotations=ro, structured_output=False)
    def bb_get_markup(path: str, markup_id: str) -> str:
        """Everything about one markup: full detail, its raw PDF dictionary keys (trimmed), grouped
        children and replies. markup_id: the id from bb_list_markups (/NM GUID, 'xref:N' or a bare xref).
        Returns {file, markup, raw_keys, group_children[], replies[]}."""
        return _j(get_markup(path, markup_id))

    @mcp.tool(annotations=ro, structured_output=False)
    def bb_takeoff_summary(path: str, group_by: Union[list[str], str, None] = None, pages: Pages = None,
                           layer: Optional[str] = None, subject: Optional[str] = None) -> str:
        """Takeoff quantities from Bluebeam measurement markups: areas (sf), lengths/perimeters (ft),
        counts and volumes (area x depth), grouped and totalled. Read-only.

        group_by: any of subject (default), layer, page, label, author, color (outline), fill_color, depth, kind, or a custom column
        name; e.g. ["subject", "layer"]. pages/layer/subject narrow the markups considered (1-based pages).
        Quantities are Bluebeam's own values (the text Revu shows, so cutouts and arcs are right) whenever
        readable in the measurement's unit; otherwise recomputed from geometry and the /Measure scale. The
        geometry is also used as a cross-check. A count takeoff's total is the number of symbols (counted per
        group, never per symbol).
        Returns {file, group_by, rows[{group, kind, markups, total, unit, pages, ids}], totals[{kind, unit,
        markups, total}], scales[{page, scale, units_per_point, markups}], warnings[], measured, skipped,
        notes?}. markups = how many markups make up the row; total = the quantity. Volume rows carry
        volume_unverified: /Depth is assumed to be in its /DepthUnit, not yet checked against Revu.
        Warnings flag unscaled measurements, several scales on one page, and Bluebeam's value differing
        from the geometry by over 0.5% (the Bluebeam value is the one used)."""
        return _j(takeoff_summary(path, group_by, pages, layer, subject))

    @mcp.tool(annotations=export_ann, structured_output=False)
    def bb_export_markups(path: str, output_path: str, format: str = "csv", pages: Pages = None,
                          overwrite: bool = False, group_by: Union[list[str], str, None] = None) -> str:
        """Export the markups list to a spreadsheet-friendly file (the PDF is not touched).

        format: "csv", "xlsx" (Markups sheet with header, frozen row, filter, plus a Takeoff sheet) or "json".
        output_path must end in the same extension and its folder must exist; an existing file is refused
        unless overwrite=True. Columns follow Revu's Markups List (Subject, Page, Page Label, Author, Date,
        Status, Layer, Comments, Measurement, Unit, Length, Area, Volume, Volume Unit, Count, Depth, Label,
        Color, Fill Color, Type, Id, Reply To) plus custom columns; replies are included. Count symbols after
        the first of a group leave the measurement cells blank so column sums stay right. group_by sets the
        Takeoff sheet's grouping (same fields as bb_takeoff_summary; default subject).
        Returns {output_path, format, rows, columns, sheets}."""
        return _j(export_markups(path, output_path, format, pages, overwrite, group_by))

    @mcp.tool(annotations=ro, structured_output=False)
    def bb_search_text(path: str, query: str, pages: Pages = None, case_sensitive: bool = False,
                       max_hits: int = 200, regex: bool = False) -> str:
        """Search the PDF's text layer (sheet numbers, notes, callouts). Read-only.

        Plain search is a phrase match (case-insensitive unless case_sensitive). regex=True matches a
        Python regex against single words only, with approximate word rects. pages is 1-based.
        Returns {file, query, total_hits, returned, truncated?, hits[{page, page_label?, rect [x0,y0,x1,y1] in
        PDF points top-left, context}], note?}. Scanned pages have no text layer and cannot be searched."""
        return _j(search_text(path, query, pages, case_sensitive, max_hits, regex))

    @mcp.tool(annotations=ro, structured_output=False)
    def bb_get_page_text(path: str, page: int, mode: str = "text", clip: Optional[list[float]] = None,
                         max_chars: int = 20000) -> str:
        """Text of one page (1-based). mode "text" = plain reading text; "blocks" = [{rect, text}] paragraphs;
        "words" = [[x0, y0, x1, y1, word]] with positions. clip = [x0, y0, x1, y1] in PDF points (origin
        top-left) limits it to a region. Output is cut at max_chars (200 to 200000) with truncated=true.
        Returns {file, page, page_label?, mode, text|blocks|words, truncated?, note?}."""
        return _j(get_page_text(path, page, mode, clip, max_chars))

    @mcp.tool(annotations=ro, structured_output=False)
    def bb_sheet_index(path: str) -> str:
        """One line per page for orienting in a drawing set: page number, page label, bookmark title,
        size in inches, orientation, rotation (if any), drawing scales (from viewports and measurement
        markups), a title-block guess (largest text in the bottom-right quarter; usually sheet number and
        title), and the markup count. Returns {file, page_count, markups, pages[]}."""
        return _j(sheet_index(path))

    @mcp.tool(annotations=ro, structured_output=False)
    def bb_render_page(path: str, page: int, max_px: int = 1600, clip: Optional[list[float]] = None,
                       annots: bool = True) -> list:
        """Render a page (1-based) or a region of it as a PNG image you can look at. Read-only.

        max_px is the longest side in pixels (64 to 2400, default 1600). clip = [x0, y0, x1, y1] in PDF
        points, origin top-left, to zoom into a region (a small clip renders sharper, up to 600 dpi).
        annots=False hides markups. The accompanying JSON gives width_px, height_px, dpi, px_per_pt and the
        clip drawn. Unrotated page: point = (clip[0] + px / px_per_pt, clip[1] + py / px_per_pt). ROTATED
        page (JSON has rotation, clip_image, px_to_pt): the image is drawn rotated while markup and text
        coordinates stay unrotated, so map pixels with px_to_pt [a,b,c,d,e,f]: x = a*px + c*py + e,
        y = b*px + d*py + f."""
        return render_page_tool(path, page, max_px, clip, annots)

    @mcp.tool(annotations=ro, structured_output=False)
    def bb_render_markup(path: str, markup_id: str, margin: float = 36, dpi: float = 150) -> list:
        """Render the area around one markup (and its grouped children) as a PNG, with margin points
        of context (default 36 = half an inch). dpi is lowered automatically to stay within 2400 px.
        markup_id from bb_list_markups. Also returns JSON with pixel size, dpi and the clip drawn (on a
        rotated page also rotation, clip_image and px_to_pt, as for bb_render_page)."""
        return render_markup_tool(path, markup_id, margin, dpi)

    @mcp.tool(annotations=ro, structured_output=False)
    def bb_compare_markups(path_a: str, path_b: str) -> str:
        """Compare the markups of two PDFs (e.g. two revisions or two reviewers' copies), matched by their
        /NM id (else page + type + position). Reports added (in B only), removed (in A only) and changed
        markups (contents, subject, author, layer, review status, measurement value, moved more than 1 pt),
        replies included. Returns {a, b, counts, added[], removed[], changed[{id, page, subject, changes}]};
        lists are capped at 50 with truncated=true."""
        return _j(compare_markups(path_a, path_b))

    @mcp.tool(annotations=ro, structured_output=False)
    def bb_list_layers(path: str) -> str:
        """Layers (optional content groups) of a PDF: name, xref, visible (default state), locked and how
        many markups sit on each. Returns {file, count, layers[]}."""
        return _j(list_layers(path))
